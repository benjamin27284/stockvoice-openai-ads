"""Add paused "function-word hint" ad groups (V6 sheet 01 column S) to the three live StockVoice campaigns.

For each campaign and each of g01, g02, g04, g06, g08, g10:
  - new ad group <name>_func, status paused, context hint = column S
  - ads copied from that campaign's live ad group (title, body, image, target URL), status paused
  - UTMs copied from live; only utm_term changes to the new ad group name
Existing ad groups and ads are not touched.

Usage:
  python push_sv_v6_func_hint_groups.py          # dry run
  python push_sv_v6_func_hint_groups.py --push

Never calls any /activate endpoint.
"""
import csv, json, sys, time
from pathlib import Path
import openpyxl, requests

sys.stdout.reconfigure(encoding="utf-8")
HERE = Path(__file__).parent
XLSX = HERE / "StockVoice_OpenAI_Ads_V6_一步距离13组_含功能词hint_2026-10-20.xlsx"
BASE = "https://api.ads.openai.com/v1"
RUN = "r1"
CAMPAIGNS = ["cmpn_36b27a64ff4881999bb5b283b33e94db", "cmpn_e289538f381081a295fb6b7058b6af4b", "cmpn_70810647a0e4819ca9c4cc07260c8dd8"]
GROUPS = ["g01_buy_after_rally", "g02_hold_or_cut", "g04_take_profit", "g06_everyone_talking",
          "g08_before_after_earnings", "g10_market_panic_or_heat"]
SUFFIX = "_func"
PLAN_FILE = HERE / "sv_v6_func_plan.json"
RESULT_FILE = HERE / "sv_v6_func_result.json"
MAPPING_CSV = HERE / "sv_v6_func_ad_id_mapping.csv"

s = requests.Session()
s.headers["Authorization"] = "Bearer " + (HERE / "ads-manager-api-key.txt").read_text().strip()


def get(path, **params):
    r = s.get(f"{BASE}/{path}", params=params, timeout=60); r.raise_for_status(); return r.json()


def post(path, body, idem):
    r = s.post(f"{BASE}/{path}", json=body, headers={"Idempotency-Key": idem}, timeout=60)
    if r.status_code != 200:
        raise RuntimeError(f"POST /{path} {r.status_code}: {r.text[:800]}")
    return r.json()


def retag_term(qs, old, new):
    parts = qs.split("&")
    if parts.count(f"utm_term={old}") != 1:
        raise ValueError(f"utm_term={old} not found once in {qs}")
    return "&".join(f"utm_term={new}" if p == f"utm_term={old}" else p for p in parts)


def build_plan():
    ws = openpyxl.load_workbook(XLSX, data_only=True)["01_组总览"]
    head = [c.value for c in ws[1]]
    assert head[18].startswith("Context hint（含功能词版"), head[18]  # column S
    hints = {r[4]: r[18] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    plan, errors = [], []
    for cid in CAMPAIGNS:
        c = get(f"campaigns/{cid}")
        live = {g["name"]: g for g in get("ad_groups", campaign_id=cid, limit=100)["data"]}
        if any(n.endswith(SUFFIX) for n in live):
            errors.append(f"{c['name']}: already has {SUFFIX} ad groups")
        pc = {"campaign_id": cid, "campaign": c["name"], "ad_groups": []}
        for name in GROUPS:
            g = live.get(name)
            if not g or not hints.get(name):
                errors.append(f"{c['name']} {name}: live group or S hint missing"); continue
            new = name + SUFFIX
            ads = sorted(get("ads", ad_group_id=g["id"], limit=100)["data"], key=lambda a: a["name"])
            if len(ads) != 2:
                errors.append(f"{c['name']} {name}: {len(ads)} live ads")
            pc["ad_groups"].append({
                "source_id": g["id"], "name": new, "context_hints": [hints[name]],
                "bidding_config": g["bidding_config"],
                "query_string_template": retag_term(g["landing_page_configuration"]["query_string_template"], name, new),
                "ads": [{
                    "source_id": a["id"], "name": a["name"].replace("_CLICK_V5", "_CLICK_V6_FUNC"),
                    "creative": {k: a["creative"][k] for k in ("type", "title", "body", "target_url", "file_id", "image_crop") if a["creative"].get(k) is not None},
                    "query_string_template": retag_term(a["landing_page_configuration"]["query_string_template"], name, new),
                } for a in ads],
            })
        plan.append(pc)
    return plan, errors


def push(plan):
    result = []
    try:
        for pc in plan:
            rc = {"campaign_id": pc["campaign_id"], "campaign": pc["campaign"], "ad_groups": []}
            result.append(rc)
            for g in pc["ad_groups"]:
                ag = post("ad_groups", {
                    "campaign_id": pc["campaign_id"], "name": g["name"], "status": "paused",
                    "context_hints": g["context_hints"], "bidding_config": g["bidding_config"],
                    "landing_page_configuration": {"query_string_template": g["query_string_template"]},
                }, f"sv-v6-func-{RUN}-{pc['campaign_id']}-{g['name']}")
                rg = {"id": ag["id"], "name": g["name"], "status": ag.get("status"), "ads": []}
                rc["ad_groups"].append(rg)
                print(f"{pc['campaign']}  ad_group {ag['id']} {g['name']} {ag.get('status')}")
                for a in g["ads"]:
                    ad = post("ads", {
                        "ad_group_id": ag["id"], "name": a["name"], "status": "paused", "creative": a["creative"],
                        "landing_page_configuration": {"query_string_template": a["query_string_template"]},
                    }, f"sv-v6-func-{RUN}-{pc['campaign_id']}-{a['name']}")
                    rg["ads"].append({"id": ad["id"], "name": a["name"], "source_id": a["source_id"], "title": a["creative"]["title"], "status": ad.get("status")})
                    print(f"    ad {ad['id']} {a['name']} {ad.get('status')}")
                    time.sleep(0.15)
    finally:
        RESULT_FILE.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        with MAPPING_CSV.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["campaign", "campaign_id", "ad_group", "ad_group_id", "ad_name", "ad_id", "source_ad_id", "title"])
            for c in result:
                for g in c["ad_groups"]:
                    for a in g["ads"]:
                        w.writerow([c["campaign"], c["campaign_id"], g["name"], g["id"], a["name"], a["id"], a["source_id"], a["title"]])


if __name__ == "__main__":
    plan, errors = build_plan()
    PLAN_FILE.write_text(json.dumps(plan, ensure_ascii=False, indent=1), encoding="utf-8")
    for pc in plan:
        print(f"{pc['campaign']}: {len(pc['ad_groups'])} ad groups / {sum(len(g['ads']) for g in pc['ad_groups'])} ads")
    if errors:
        sys.exit("errors:\n  " + "\n  ".join(errors))
    if "--push" in sys.argv:
        push(plan)
