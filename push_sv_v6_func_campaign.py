"""Create one paused "function-word hint" campaign per market, paired with each live StockVoice campaign.

For each source campaign (TW = STOCKVOICE_US_CONTEXT_CLICK_V5, US = ..._V5_US, CA = ..._V5_CA):
  - new campaign with the same live settings (countries, budget, conversion events), status paused
  - ad groups g01/g02/g04/g06/g08/g10 + "_func", context hint = V6 sheet 01 column S
  - ads copied from the source's live ads (title, body, image, target URL)
  - UTMs copied from live; utm_campaign gets "_func", utm_term gets "_func"

Usage:
  python push_sv_v6_func_campaign.py          # dry run
  python push_sv_v6_func_campaign.py --push

Never calls any /activate endpoint.
"""
import csv, json, sys, time
from pathlib import Path
import openpyxl, requests

sys.stdout.reconfigure(encoding="utf-8")
HERE = Path(__file__).parent
XLSX = HERE / "StockVoice_OpenAI_Ads_V6_一步距离13组_含功能词hint_2026-10-20.xlsx"
BASE = "https://api.ads.openai.com/v1"
RUN = "r1"
SOURCES = {  # source campaign id -> new campaign name
    "cmpn_36b27a64ff4881999bb5b283b33e94db": "STOCKVOICE_TW_CONTEXT_CLICK_V6_FUNC",
    "cmpn_e289538f381081a295fb6b7058b6af4b": "STOCKVOICE_US_CONTEXT_CLICK_V6_FUNC_US",
    "cmpn_70810647a0e4819ca9c4cc07260c8dd8": "STOCKVOICE_CA_CONTEXT_CLICK_V6_FUNC_CA",
}
GROUPS = ["g01_buy_after_rally", "g02_hold_or_cut", "g04_take_profit", "g06_everyone_talking",
          "g08_before_after_earnings", "g10_market_panic_or_heat"]
SUFFIX = "_func"
PLAN_FILE = HERE / "sv_v6_func_campaign_plan.json"
RESULT_FILE = HERE / "sv_v6_func_campaign_result.json"
MAPPING_CSV = HERE / "sv_v6_func_campaign_ad_id_mapping.csv"

s = requests.Session()
s.headers["Authorization"] = "Bearer " + (HERE / "ads-manager-api-key.txt").read_text().strip()


def get(path, **params):
    r = s.get(f"{BASE}/{path}", params=params, timeout=60); r.raise_for_status(); return r.json()


def post(path, body, idem):
    r = s.post(f"{BASE}/{path}", json=body, headers={"Idempotency-Key": idem}, timeout=60)
    if r.status_code != 200:
        raise RuntimeError(f"POST /{path} {r.status_code}: {r.text[:800]}")
    return r.json()


def retag(qs, term=None):
    """Append SUFFIX to utm_campaign (and to utm_term if given); keep every other param and the order."""
    out, hits = [], 0
    for p in qs.split("&"):
        k, _, v = p.partition("=")
        if k == "utm_campaign":
            p, hits = f"{k}={v}{SUFFIX}", hits + 1
        elif k == "utm_term" and term:
            if v != term:
                raise ValueError(f"utm_term {v} != {term}")
            p, hits = f"{k}={v}{SUFFIX}", hits + 1
        out.append(p)
    if hits != (2 if term else 1):
        raise ValueError(f"unexpected UTM params in {qs}")
    return "&".join(out)


def build_plan():
    ws = openpyxl.load_workbook(XLSX, data_only=True)["01_组总览"]
    assert ws.cell(1, 19).value.startswith("Context hint（含功能词版"), ws.cell(1, 19).value  # column S
    hints = {r[4]: r[18] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    existing = {c["name"] for c in get("campaigns", limit=100)["data"] if c["status"] != "archived"}
    plan, errors = [], []
    for cid, new_name in SOURCES.items():
        src = get(f"campaigns/{cid}")
        if new_name in existing:
            errors.append(f"{new_name} already exists")
        live = {g["name"]: g for g in get("ad_groups", campaign_id=cid, limit=100)["data"] if g["status"] != "archived"}
        pc = {"source_id": cid, "source": src["name"], "campaign": {
            "name": new_name, "status": "paused", "objective": src["objective"], "bidding_type": src["bidding_type"],
            "billing_event_type": src["billing_event_type"], "conversion_event_setting_ids": src["conversion_event_setting_ids"],
            "budget": src["budget"],
            "targeting": {"locations": {"countries": [l["country_code"] for l in src["targeting"]["locations"]["include"]]}},
            "landing_page_configuration": {"query_string_template": retag(src["landing_page_configuration"]["query_string_template"])},
        }, "ad_groups": []}
        for name in GROUPS:
            g = live.get(name)
            if not g or not hints.get(name):
                errors.append(f"{src['name']} {name}: live group or S hint missing"); continue
            ads = sorted((a for a in get("ads", ad_group_id=g["id"], limit=100)["data"] if a["status"] != "archived"), key=lambda a: a["name"])
            if len(ads) != 2:
                errors.append(f"{src['name']} {name}: {len(ads)} live ads")
            pc["ad_groups"].append({
                "source_id": g["id"], "name": name + SUFFIX, "context_hints": [hints[name]], "bidding_config": g["bidding_config"],
                "query_string_template": retag(g["landing_page_configuration"]["query_string_template"], name),
                "ads": [{"source_id": a["id"], "name": a["name"].replace("_CLICK_V5", "_CLICK_V6_FUNC"),
                         "creative": {k: a["creative"][k] for k in ("type", "title", "body", "target_url", "file_id", "image_crop") if a["creative"].get(k) is not None},
                         "query_string_template": retag(a["landing_page_configuration"]["query_string_template"], name)} for a in ads],
            })
        plan.append(pc)
    return plan, errors


def push(plan):
    result = []
    try:
        for pc in plan:
            name = pc["campaign"]["name"]
            camp = post("campaigns", pc["campaign"], f"{name}-{RUN}-campaign")
            rc = {"id": camp["id"], "name": name, "status": camp.get("status"), "source_id": pc["source_id"], "ad_groups": []}
            result.append(rc)
            print(f"campaign {camp['id']} {name} {camp.get('status')}")
            for g in pc["ad_groups"]:
                ag = post("ad_groups", {
                    "campaign_id": camp["id"], "name": g["name"], "status": "paused",
                    "context_hints": g["context_hints"], "bidding_config": g["bidding_config"],
                    "landing_page_configuration": {"query_string_template": g["query_string_template"]},
                }, f"{name}-{RUN}-{g['name']}")
                rg = {"id": ag["id"], "name": g["name"], "ads": []}
                rc["ad_groups"].append(rg)
                print(f"  ad_group {ag['id']} {g['name']} {ag.get('status')}")
                for a in g["ads"]:
                    ad = post("ads", {
                        "ad_group_id": ag["id"], "name": a["name"], "status": "paused", "creative": a["creative"],
                        "landing_page_configuration": {"query_string_template": a["query_string_template"]},
                    }, f"{name}-{RUN}-{a['name']}")
                    rg["ads"].append({"id": ad["id"], "name": a["name"], "source_id": a["source_id"], "title": a["creative"]["title"]})
                    print(f"    ad {ad['id']} {a['name']} {ad.get('status')}")
                    time.sleep(0.15)
    finally:
        RESULT_FILE.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        with MAPPING_CSV.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["campaign", "campaign_id", "ad_group", "ad_group_id", "ad_name", "ad_id", "source_ad_id", "title"])
            for c in result:
                for g in c["ad_groups"]:
                    for a in g["ads"]:
                        w.writerow([c["name"], c["id"], g["name"], g["id"], a["name"], a["id"], a["source_id"], a["title"]])


if __name__ == "__main__":
    plan, errors = build_plan()
    PLAN_FILE.write_text(json.dumps(plan, ensure_ascii=False, indent=1), encoding="utf-8")
    for pc in plan:
        c = pc["campaign"]
        print(f"{pc['source']} -> {c['name']}: {c['targeting']['locations']['countries']} / "
              f"TWD {c['budget'].get('daily_spend_limit_micros', 0) // 1_000_000}/day / {len(c['conversion_event_setting_ids'])} events / "
              f"{len(pc['ad_groups'])} groups, {sum(len(g['ads']) for g in pc['ad_groups'])} ads")
        print(f"    {c['landing_page_configuration']['query_string_template']}")
        g = pc["ad_groups"][0]
        print(f"    e.g. {g['name']}: {g['query_string_template']}")
    if errors:
        sys.exit("errors:\n  " + "\n  ".join(errors))
    if "--push" in sys.argv:
        push(plan)
